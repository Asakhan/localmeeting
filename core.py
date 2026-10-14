"""Local Minutes: durable storage, canonical speaker IDs, Word export."""
from pathlib import Path
from datetime import datetime
import json, os, re, sys, uuid

ROOT = Path(sys.executable).parent if getattr(sys, 'frozen', False) else Path(__file__).resolve().parent
DATA = Path(os.environ.get('LOCALAPPDATA', str(Path.home()))) / 'LocalMinutes'
DEFAULT_FOLDER = Path.home() / 'Documents' / '회의록'
SECTIONS = [('agenda', '안건 (Agenda)'), ('discussion', '주요 논의사항'),
            ('suggestions', '제안 사항'), ('decisions', '결정 사항'),
            ('unresolved', '미결 사항'), ('actions', '실행 과제 (Action Items)'),
            ('follow_up', '후속 일정')]

def now():
    return datetime.now().astimezone().isoformat(timespec='seconds')

def atomic_json(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + '.' + uuid.uuid4().hex + '.tmp')
    with tmp.open('w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.flush(); os.fsync(f.fileno())
    os.replace(tmp, path)

def read_json(path, default=None):
    try:
        return json.loads(Path(path).read_text(encoding='utf-8-sig'))
    except (OSError, ValueError):
        return default

def settings():
    defaults = dict(folder=str(DEFAULT_FOLDER), language='ko', system_audio=False,
                    microphone=-1, threads=4, speaker_threshold=0.65, glossary='',
                    device='auto', model_root=str(ROOT / 'models'))
    defaults.update(read_json(DATA / 'settings.json', {}) or {})
    return defaults

def stamp(seconds):
    seconds = max(0, int(seconds))
    return f'{seconds//3600:02}:{seconds//60%60:02}:{seconds%60:02}'

def display(value, names):
    # Render from canonical IDs every time; never chained replace (1 vs 10).
    if isinstance(value, str):
        return re.sub(r'참석자\d+(?!\d)', lambda m: names.get(m[0], m[0]), value)
    if isinstance(value, list):
        return [display(x, names) for x in value]
    if isinstance(value, dict):
        return {k: display(v, names) for k,v in value.items()}
    return value

def transcript(folder):
    rows = []
    path = Path(folder) / 'transcript.jsonl'
    if path.exists():
        for line in path.read_text(encoding='utf-8').splitlines():
            try: rows.append(json.loads(line))
            except ValueError: pass # trailing interrupted write, original audio preserved
    return sorted(rows, key=lambda r: (r['start'], r.get('source','')))

def append_jsonl(path, data):
    with Path(path).open('a', encoding='utf-8') as f:
        f.write(json.dumps(data, ensure_ascii=False) + '\n'); f.flush(); os.fsync(f.fileno())

def empty_minutes(reason='발언에서 확인되지 않음'):
    return dict(title='회의 기록', **{k: [] for k,_ in SECTIONS}, note=reason)

def validate_minutes(data):
    if not isinstance(data, dict) or not isinstance(data.get('title'), str):
        raise ValueError('회의록 응답 형식이 올바르지 않습니다.')
    result = {'title': data['title'][:160] or '회의 기록'}
    for key,_ in SECTIONS:
        value = data.get(key, [])
        if not isinstance(value, list):
            raise ValueError(f'회의록 {key} 항목이 목록이 아닙니다.')
        if key == 'actions':
            if any(not isinstance(x,dict) for x in value): raise ValueError('실행 과제 형식 오류')
            result[key] = [{k: str(x.get(k) or '미정') for k in ('owner','due','task')} for x in value]
        else:
            if any(not isinstance(x,str) for x in value): raise ValueError(f'{key} 내용 형식 오류')
            result[key] = value
    return result

def export_docx(folder):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    folder = Path(folder)
    meta = read_json(folder/'meeting.json', {})
    names = read_json(folder/'names.json', {})
    raw = read_json(folder/'minutes.json', empty_minutes('자동 정리가 아직 완료되지 않았습니다.'))
    minutes = display(raw, names)
    rows = transcript(folder)
    speakers = sorted({r['speaker'] for r in rows if r['speaker'].startswith('참석자')}, key=lambda x:int(x[3:]))
    doc = Document()
    sec=doc.sections[0]; sec.top_margin=Cm(2); sec.bottom_margin=Cm(2)
    normal = doc.styles['Normal']; normal.font.name='맑은 고딕'; normal.font.size=Pt(10)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    for style in ['Title','Heading 1','Heading 2']:
        doc.styles[style].font.name='맑은 고딕'
        doc.styles[style].font.color.rgb=RGBColor.from_string('163B36')
        doc.styles[style]._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    doc.add_heading(minutes.get('title','회의 기록'), 0)
    doc.add_paragraph('LOCAL MINUTES  ·  로컬 AI 회의록', style='Subtitle')
    table=doc.add_table(rows=0,cols=2); table.style='Light Shading Accent 1'
    for label,value in [('회의명', minutes.get('title','회의 기록')), ('일시', f"{meta.get('started','')} ~ {meta.get('ended','진행 중')}"),
                        ('참석자', f"음성으로 추정한 {len(speakers)}명 · "+', '.join(names.get(s,s) for s in speakers))]:
        cells=table.add_row().cells; cells[0].text=label; cells[1].text=value
    if raw.get('note'): doc.add_paragraph(raw['note'])
    for key,title in SECTIONS:
        doc.add_heading(title, level=1)
        values=minutes.get(key,[])
        if not values:
            doc.add_paragraph('발언에서 확인되지 않음'); continue
        if key == 'actions':
            t=doc.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
            for c,text in zip(t.rows[0].cells,['담당자','기한','할 일']):c.text=text
            for row in values:
                for c,k in zip(t.add_row().cells,['owner','due','task']):c.text=str(row.get(k,'미정'))
        else:
            for value in values: doc.add_paragraph(str(value),style='List Bullet')
    doc.add_page_break();doc.add_heading('발언 기록',level=1)
    doc.add_paragraph('화자와 참석 인원은 음성 기반 추정입니다. 겹침 발화·짧은 발언은 확인이 필요합니다.')
    for row in rows:
        p=doc.add_paragraph()
        p.add_run(f"[{stamp(row['start'])}] {names.get(row['speaker'],row['speaker'])}  ").bold=True
        p.add_run(display(row['text'],names))
    target=folder/'회의록.docx'
    temp=folder/'회의록.tmp.docx';doc.save(temp)
    try:os.replace(temp,target)
    except PermissionError:
        target=folder/f'회의록_수정_{datetime.now():%H%M%S}.docx';os.replace(temp,target)
    return target

def new_meeting(config):
    root=Path(config['folder']);root.mkdir(parents=True,exist_ok=True)
    folder=root/(datetime.now().strftime('%Y-%m-%d_%H%M%S')+'_'+uuid.uuid4().hex[:6])
    (folder/'audio').mkdir(parents=True)
    atomic_json(folder/'meeting.json',dict(started=now(),ended=None,state='recording',config=config))
    atomic_json(folder/'names.json',{})
    return folder

def export_text(folder):
    folder=Path(folder);names=read_json(folder/'names.json',{})
    meta=read_json(folder/'meeting.json',{})
    lines=[f"회의 시작: {meta.get('started','')}",f"회의 종료: {meta.get('ended','')}",'']
    for row in transcript(folder):
        lines.append(f"[{stamp(row['start'])} ~ {stamp(row['end'])}] {names.get(row['speaker'],row['speaker'])}: {display(row['text'],names)}")
    temp=folder/'전체전사.tmp'
    with temp.open('w',encoding='utf-8-sig') as f:
        f.write('\n'.join(lines)+'\n');f.flush();os.fsync(f.fileno())
    os.replace(temp,folder/'전체전사.txt')
    return folder/'전체전사.txt'

def finalize_outputs(folder, allow_delete=True):
    """Delete only this meeting's audio, only after all outputs validate."""
    import zipfile
    folder=Path(folder).resolve()
    doc=export_docx(folder);text=export_text(folder)
    with zipfile.ZipFile(doc) as z:
        if z.testzip() is not None or 'word/document.xml' not in z.namelist():raise RuntimeError('Word 파일 검증 실패')
    if not text.read_text(encoding='utf-8-sig').strip():raise RuntimeError('전체 전사 저장 확인 실패')
    meta=read_json(folder/'meeting.json',{})
    capture=read_json(folder/'capture_done.json',{})
    missing=[p for p in (folder/'audio').glob('*.json') if 'file' in (read_json(p,{}) or {}) and not (folder/'results'/p.name).exists()]
    can_delete=allow_delete and not missing and not capture.get('errors') and not list((folder/'audio').glob('*.part.wav'))
    if can_delete:
        audio=(folder/'audio').resolve()
        if audio.parent!=folder or audio.is_symlink():raise RuntimeError('오디오 폴더 경로 확인 실패')
        for p in audio.iterdir():
            if p.is_file() and not p.is_symlink() and p.suffix.lower() in ('.wav','.json'):
                p.unlink()
        meta['audio_deleted']=True
    else:meta['audio_deleted']=False
    meta['docx']=doc.name;meta['state']='complete' if can_delete else 'review'
    atomic_json(folder/'meeting.json',meta)
    return doc
